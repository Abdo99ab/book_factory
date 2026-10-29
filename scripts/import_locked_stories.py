#!/usr/bin/env python3
"""Parse APPROVED / FINAL CHECK story .docx files into brand/stories/<title_id>.json."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
STORIES = ROOT / "brand" / "stories"
CATALOGUE = ROOT / "brand" / "catalogue.json"

DOC_MAP = {
    "01": "dinosaur_egg",
    "02": "big_kid_list",
    "03": "nursery_day",
    "04": "bed_boat",
    "05": "night_rocket",
    "06": "ocean_map",
    "07": "tells_truth",
    "08": "worry_jar",
    "09": "new_front_door",
    "10": "too_good_picture",
}

READ_AGAIN_NIGHTS = (
    "Read this story a few nights in a row, and once more just before the day itself."
)
READ_AGAIN_ASKS = "Read it a few nights in a row, and again whenever [Name] asks."
READ_AGAIN_DAY = {"nursery_day", "bed_boat", "new_front_door"}

# Recurring characters/props seeded from each title's locked Anchor (sheets unapproved until generated).
CHARACTER_SEEDS: dict[str, list[dict]] = {
    "ocean_map": [
        {
            "id": "tide",
            "name": "Tide",
            "description": (
                "Tide the turtle guide: shell like a green umbrella, kind eyes, rounded flippers. "
                "Child-scale companion (not tiny). Neutral swimming/standing pose. "
                "Front view and side view on a plain blank background. No child, no map, no letters."
            ),
            "sheet": "brand/stories/ocean_map/tide.png",
            "approved": False,
            "spreads": list(range(2, 14)),
        },
        {
            "id": "sea_glass_compass",
            "name": "Sea-glass compass",
            "kind": "prop",
            "aliases": ["compass", "sea-glass compass"],
            "description": (
                "Treasure compass made of sea glass: soft teal and oat glass, simple needle, "
                "fits in a child's palm. Two clear views on a plain blank background. "
                "No child, no turtle, no letters."
            ),
            "sheet": "brand/stories/ocean_map/sea_glass_compass.png",
            "approved": False,
            "spreads": [10, 11, 12, 13, 14],
        },
    ],
    "big_kid_list": [
        {
            "id": "fridge_list",
            "name": "Fridge list",
            "kind": "prop",
            "aliases": ["list", "big-kid list", "fridge list"],
            "description": (
                "Paper list on a fridge with five simple picture icons: shoes, coat, cup, hands, potty. "
                "Blank ticks (no ticks filled). Two clear views on a plain blank background. No child, no letters beyond simple icons."
            ),
            "sheet": "brand/stories/big_kid_list/fridge_list.png",
            "approved": False,
            "spreads": list(range(1, 15)),
        },
    ],
    "night_rocket": [
        {
            "id": "night_rocket",
            "name": "Night rocket",
            "kind": "prop",
            "aliases": ["rocket", "night rocket"],
            "description": (
                "Friendly cardboard-and-paint night rocket for a child's bedroom adventure: "
                "rounded, soft colours (oat, teal accents), no sharp fins, visor never covering a face. "
                "Two clear views on a plain blank background. No child, no logos, no letters."
            ),
            "sheet": "brand/stories/night_rocket/night_rocket.png",
            "approved": False,
            "spreads": list(range(1, 15)),
        },
    ],
    "worry_jar": [
        {
            "id": "worry_jar",
            "name": "Worry jar",
            "kind": "prop",
            "aliases": ["jar", "worry jar"],
            "description": (
                "Clear glass worry jar with a simple lid, soft oat/butter paper slips inside optional. "
                "Two clear views on a plain blank background. No child, no scary faces, no letters."
            ),
            "sheet": "brand/stories/worry_jar/worry_jar.png",
            "approved": False,
            "spreads": list(range(1, 15)),
        },
    ],
    "too_good_picture": [
        {
            "id": "picture",
            "name": "Too-good picture",
            "kind": "prop",
            "aliases": ["picture", "photo", "screen"],
            "description": (
                "A bright too-good-to-be-true picture / tablet screen showing a perfect scene "
                "(generic, no brand logos). Two clear views on a plain blank background. No child."
            ),
            "sheet": "brand/stories/too_good_picture/picture.png",
            "approved": False,
            "spreads": list(range(1, 15)),
        },
    ],
}


def _paras(doc: Document) -> list[str]:
    return [p.text.strip() for p in doc.paragraphs]


def _table_map(table) -> dict[str, str]:
    out = {}
    for row in table.rows[1:]:
        if len(row.cells) < 2:
            continue
        key = row.cells[0].text.strip().lower()
        val = row.cells[1].text.strip()
        if key and val:
            out[key] = val
    return out


def _pick(mapping: dict[str, str], *keys: str) -> str:
    for key in keys:
        for mk, mv in mapping.items():
            if key in mk:
                return mv
    return ""


def _split_art_meta(blob: str) -> tuple[str, str, str]:
    """Split a combined Art/Anchor/Calm/Route string into (art, anchor, calm_note)."""
    text = (blob or "").strip()
    if not text:
        return "", "", ""
    art = text
    anchor = ""
    calm_note = ""
    m = re.search(r"\bAnchor:\s*", text, re.I)
    if m:
        art = text[: m.start()].strip().rstrip(" .;")
        rest = text[m.end() :]
        calm_m = re.search(r"\bCalm zone:\s*", rest, re.I)
        if calm_m:
            anchor = rest[: calm_m.start()].strip().rstrip(" .;")
            rest2 = rest[calm_m.end() :]
            route_m = re.search(r"\bRoute:\s*", rest2, re.I)
            calm_note = (rest2[: route_m.start()] if route_m else rest2).strip().rstrip(" .;")
        else:
            route_m = re.search(r"\bRoute:\s*", rest, re.I)
            anchor = (rest[: route_m.start()] if route_m else rest).strip().rstrip(" .;")
    else:
        # Still strip trailing Route: if present without Anchor
        route_m = re.search(r"\bRoute:\s*", text, re.I)
        calm_m = re.search(r"\bCalm zone:\s*", text, re.I)
        if calm_m:
            art = text[: calm_m.start()].strip().rstrip(" .;")
            rest2 = text[calm_m.end() :]
            route_m = re.search(r"\bRoute:\s*", rest2, re.I)
            calm_note = (rest2[: route_m.start()] if route_m else rest2).strip().rstrip(" .;")
        elif route_m:
            art = text[: route_m.start()].strip().rstrip(" .;")
    return art, anchor, calm_note


def _calm_side(note: str, n: int) -> str:
    low = (note or "").lower()
    # Prefer explicit left/right-hand / page cues.
    if re.search(r"right[- ]hand|right page|top right|lower right|upper right|\bright\b", low):
        if re.search(r"left[- ]hand|left page|top left|lower left|upper left", low):
            # Both mentioned — take the last side word.
            last_left = max(
                (m.end() for m in re.finditer(r"left[- ]hand|left page|\bleft\b", low)),
                default=-1,
            )
            last_right = max(
                (m.end() for m in re.finditer(r"right[- ]hand|right page|\bright\b", low)),
                default=-1,
            )
            return "right" if last_right >= last_left else "left"
        return "right"
    if re.search(r"left[- ]hand|left page|top left|lower left|upper left|\bleft\b", low):
        return "left"
    # Sky / upper / stars without side → alternate starting right on spread 1.
    return "right" if int(n) % 2 == 1 else "left"


def _parse_spread_block(lines: list[str], start: int) -> tuple[dict, int]:
    header = lines[start]
    m = re.match(r"Spread\s+(\d+)", header, re.I)
    if not m:
        raise ValueError(f"Not a spread header: {header}")
    n = int(m.group(1))
    text = ""
    art_raw = ""
    anchor_raw = ""
    i = start + 1
    while i < len(lines):
        line = lines[i]
        if re.match(r"Spread\s+\d+", line, re.I) or line.startswith("Pp ") or line.startswith("P 32") or line.startswith("Pages 30"):
            break
        if line.startswith("Art:"):
            art_raw = line[4:].strip()
        elif line.startswith("Anchor:"):
            anchor_raw = line[7:].strip()
        elif not text and line and not re.match(r"^\d+\s*words", line, re.I):
            if not line.startswith("Art:") and not line.startswith("Anchor:"):
                text = line
        i += 1

    # Prefer separate Anchor line; else strip meta from Art line.
    if anchor_raw:
        art, _, _ = _split_art_meta(art_raw)
        _, anchor, calm_note = _split_art_meta("Anchor: " + anchor_raw)
        if not calm_note:
            # Calm may still be on the art line in some drafts
            _, _, calm_note = _split_art_meta(art_raw)
            if art_raw and not art:
                art, _, _ = _split_art_meta(art_raw)
    else:
        art, anchor, calm_note = _split_art_meta(art_raw)

    setting = art
    if ";" in art:
        setting = art.split(";", 1)[0].strip()
    # Drop leading lighting-only fragments that leave setting empty
    if not setting:
        setting = art

    return {
        "n": n,
        "text": text,
        "setting": setting,
        "art": art,
        "anchor": anchor,
        "calm_zone": _calm_side(calm_note, n),
        "calm_zone_note": (f"{calm_note} — keep low-detail; story text sits there.") if calm_note else "",
    }, i


def _family_meta(title_id: str, catalogue: dict) -> dict:
    return (catalogue.get("titles") or {}).get(title_id) or {}


def _must_not_list(raw: str, fallback: list) -> list:
    if not raw:
        return list(fallback or [])
    # Prefer comma / semicolon splits; strip leading "No ".
    parts = re.split(r"[;•\n]|(?<=\.),\s+(?=[A-Z])", raw)
    out = []
    for p in parts:
        s = p.strip().rstrip(".")
        if s.lower().startswith("no "):
            s = s[3:].strip()
        if s:
            out.append(s)
    return out or list(fallback or [])


def parse_story_doc(path: Path, title_id: str, catalogue: dict) -> dict:
    doc = Document(str(path))
    lines = [t for t in _paras(doc) if t]
    cover = _table_map(doc.tables[0]) if doc.tables else {}
    parent_map = {}
    for table in doc.tables[1:]:
        header = table.rows[0].cells[0].text.strip().lower() if table.rows else ""
        if "block" in header:
            parent_map = _table_map(table)
            break

    cat = _family_meta(title_id, catalogue)
    band = _pick(cover, "front band") or cat.get("band") or ""
    title_tmpl = ""
    for k, v in cover.items():
        if k.startswith("title"):
            title_tmpl = v
            break
    title_tmpl = title_tmpl or cat.get("title_template") or ""
    question = _pick(cover, "back-cover question", "question") or cat.get("question_form") or ""
    blurb = _pick(cover, "blurb", "back-cover blurb") or cat.get("blurb") or ""
    honesty = _pick(cover, "fixed line", "back-cover fixed") or (
        "Illustrated from [Name]'s photo and checked by a person before print."
    )
    must_not = _must_not_list(_pick(cover, "must not"), list(cat.get("must_not_show") or []))
    anchor_desc = _pick(cover, "anchor") or (cat.get("anchor") or {}).get("description") or ""

    outfit = cat.get("outfit") or ""
    cover_art = ""
    route = cat.get("route") or "morning_light"
    for line in lines:
        low = line.lower()
        if low.startswith("cover art:"):
            cover_art = line.split(":", 1)[1].strip()
        if "outfit" in low and ("jumper" in low or "trousers" in low or "pyjamas" in low or "swimsuit" in low or "rash" in low or "fixed:" in low):
            m = re.search(r"(?:outfit|swimwear|fixed)[^:]*:\s*(.+)", line, re.I)
            if m:
                outfit = m.group(1).split(".")[0].strip()
        if line.startswith("Route:") or "route:" in low[:20]:
            if "Field Guide" in line:
                route = "field_guide"
            elif "Morning Light" in line:
                route = "morning_light"

    # Cover-table swimwear / outfit overrides
    if "rash" in (anchor_desc or "").lower() or "swim" in (outfit or "").lower():
        if "rash vest" in anchor_desc.lower():
            outfit = "rash vest and swim shorts"

    spot = ""
    for i, line in enumerate(lines):
        if "Spot illustration" in line or line.startswith("Spot:"):
            spot = re.sub(r"^Spot(?: illustration)?[^:]*:\s*", "", line, flags=re.I).strip()
        if line.startswith("Page 1") or line.startswith("P 1") or "Page 1:" in line:
            # e.g. "Page 1: band, title, … small spot of the bottle."
            if "spot" in line.lower():
                spot = re.sub(r"^.*\bspot(?: of|:)?\s*", "", line, flags=re.I).strip().rstrip(".")
                spot = f"Small spot of {spot}. No child on this page." if not spot.lower().startswith("small") else spot

    spreads: list[dict] = []
    i = 0
    while i < len(lines):
        if re.match(r"Spread\s+\d+", lines[i], re.I):
            sp, i = _parse_spread_block(lines, i)
            if sp.get("text") and sp.get("art"):
                spreads.append(sp)
            continue
        i += 1

    if len(spreads) != 14:
        raise ValueError(f"{path.name}: expected 14 spreads, got {len(spreads)}")

    for sp in spreads:
        if not sp["setting"]:
            sp["setting"] = sp["art"]

    # Fallback cover art from spread 1 / catalogue if missing
    if not cover_art and spreads:
        cover_art = (
            f"[Name] in the story's key scene from spread 1 — face large and front-on; "
            f"upper third calm for the title. {spreads[0]['art'][:180]}"
        )

    parent = {
        "you_could_say": _pick(parent_map, "you could say")
        or (cat.get("parent_spread") or {}).get("you_could_say")
        or "",
        "try_together": _pick(parent_map, "try together")
        or (cat.get("parent_spread") or {}).get("try_together")
        or "",
        "in_real_life": _pick(parent_map, "in real life")
        or (cat.get("parent_spread") or {}).get("in_real_life")
        or "",
        "read_it_again": READ_AGAIN_NIGHTS if title_id in READ_AGAIN_DAY else READ_AGAIN_ASKS,
    }

    age_band = cat.get("age_band") or "3-6"
    family = cat.get("family") or "adventures"
    colour = cat.get("colour") or "teal"

    corpus = "\n".join(
        [title_tmpl, blurb, cover_art, spot, parent["read_it_again"], anchor_desc]
        + [sp["text"] + " " + sp["art"] for sp in spreads]
    )
    order_inputs = {
        "carer": bool(re.search(r"\[Carer\]", corpus)),
        "comfort_object": bool(re.search(r"\[Object\]", corpus)),
        "giver": True,
    }

    characters = []
    for seed in CHARACTER_SEEDS.get(title_id) or []:
        entry = dict(seed)
        entry["title_id"] = title_id
        cid = entry.get("id") or "character"
        sheet = (entry.get("sheet") or "").replace("\\", "/")
        if not sheet.startswith(f"brand/stories/{title_id}/"):
            entry["sheet"] = f"brand/stories/{title_id}/{cid}.png"
        characters.append(entry)

    data = {
        "title_id": title_id,
        "approved": True,
        "version": "v2-imported",
        "status": "final",
        "source": path.name,
        "family": family,
        "colour": colour,
        "age_band": age_band,
        "route": route,
        "front_band": band,
        "title_template": title_tmpl.replace("{name}", "[Name]"),
        "back_question": question.replace("{name}", "[Name]"),
        "back_blurb": blurb,
        "honesty_line": honesty,
        "spine_title": title_tmpl.replace("{name}", "[Name]"),
        "outfit": outfit,
        "cover_art": cover_art,
        "title_page": {
            "band": band,
            "line": "Illustrated from [Name]'s photo",
            "spot": spot or "Story anchor object. No child on this page.",
        },
        "anchor": {
            "description": anchor_desc,
            "lock_spread": 2 if title_id == "dinosaur_egg" else None,
        },
        "must_not_show": must_not,
        "characters": characters,
        "personalisation_slots": cat.get("personalisation_slots") or [],
        "order_inputs": order_inputs,
        "spreads": spreads,
        "parent_spread": parent,
        "parent_art": "",
        "dedication_template": "For [Name], from [Giver]",
        "last_line": "",
    }
    if title_id == "ocean_map":
        data["inworld_lettering"] = [
            {
                "spread": 1,
                "kind": "map_banner_name",
                "font": "inworld",
                "fallback_font": "body",
                "side": "right",
                "region": {
                    # Fractions of the right page (map banner near the top).
                    "x0": 0.12,
                    "y0": 0.08,
                    "x1": 0.88,
                    "y1": 0.20,
                },
                "curve": True,
                "note": (
                    "Caveat (OFL) for the Ocean Map name on the blank banner only — "
                    "never story text, titles, bands, covers, or ads. "
                    "Test short and long names curved to the banner; if it looks pasted on, use Nunito."
                ),
            }
        ]
        # Name overlay note from cover table
        overlay = _pick(cover, "name overlay")
        if overlay:
            data["inworld_lettering"][0]["source_note"] = overlay

    t14 = spreads[-1]["text"]
    parts = re.split(r"(?<=[.!?…\"”])\s+", t14.strip())
    if parts:
        data["last_line"] = parts[-1]
    return data


def find_docs() -> dict[str, Path]:
    found = {}
    for path in STORIES.glob("*.docx"):
        m = re.match(r"^(\d{2})\b", path.name)
        if not m or m.group(1) not in DOC_MAP:
            continue
        tid = DOC_MAP[m.group(1)]
        score = 0
        if "FINAL" in path.name.upper():
            score += 2
        if "APPROVED" in path.name.upper():
            score += 1
        if tid not in found or score > found[tid][0]:
            found[tid] = (score, path)
    return {tid: path for tid, (_s, path) in found.items()}


def main() -> int:
    catalogue = json.loads(CATALOGUE.read_text(encoding="utf-8"))
    docs = find_docs()
    print("Found docs:", {k: v.name for k, v in docs.items()})
    for title_id, path in sorted(docs.items()):
        out_path = STORIES / f"{title_id}.json"
        existing_chars = []
        old = None
        if out_path.is_file():
            try:
                old = json.loads(out_path.read_text(encoding="utf-8"))
                existing_chars = list(old.get("characters") or [])
            except Exception:
                existing_chars = []
        data = parse_story_doc(path, title_id, catalogue)

        # Prefer existing character sheets/approvals when ids match.
        if existing_chars:
            by_id = {c.get("id"): c for c in existing_chars if c.get("id")}
            if title_id == "dinosaur_egg":
                data["characters"] = existing_chars
            else:
                merged = []
                used = set()
                for seed in data["characters"]:
                    prev = by_id.get(seed["id"])
                    if prev:
                        merged.append({**seed, **{k: prev[k] for k in prev if k in (
                            "approved", "approved_by", "approved_at", "sheet", "description", "name", "kind", "aliases", "spreads"
                        )}})
                        used.add(seed["id"])
                    else:
                        merged.append(seed)
                for c in existing_chars:
                    if c.get("id") not in used:
                        merged.append(c)
                data["characters"] = merged

        if title_id == "dinosaur_egg" and old:
            data["version"] = old.get("version") or data["version"]
            if "in arms" in (old.get("cover_art") or "") or "beside the child" in (old.get("cover_art") or ""):
                data["cover_art"] = old["cover_art"]
            if old.get("parent_art"):
                data["parent_art"] = old["parent_art"]
            # Keep polished self-contained settings/art when present.
            by_n = {int(s["n"]): s for s in (old.get("spreads") or [])}
            for sp in data["spreads"]:
                prev = by_n.get(int(sp["n"]))
                if not prev:
                    continue
                if prev.get("setting") and len(str(prev["setting"])) > 40:
                    sp["setting"] = prev["setting"]
                prev_art = prev.get("art") or ""
                if prev_art and "Same spot" not in prev_art and find_relative_safe(prev_art):
                    sp["art"] = prev_art
                    sp["calm_zone"] = prev.get("calm_zone") or sp["calm_zone"]
                    sp["calm_zone_note"] = prev.get("calm_zone_note") or sp["calm_zone_note"]
            if old.get("anchor") and (old["anchor"].get("description") or ""):
                # Keep detailed Pip lock if longer
                if len(str(old["anchor"].get("description") or "")) > len(str(data["anchor"].get("description") or "")):
                    data["anchor"] = old["anchor"]

        out_path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"Wrote {out_path.name} ({len(data['spreads'])} spreads, route={data['route']}, chars={len(data['characters'])})")
        if title_id in catalogue.get("titles", {}):
            catalogue["titles"][title_id]["approved"] = True
            if data.get("front_band"):
                catalogue["titles"][title_id]["band"] = data["front_band"]
            if data.get("outfit"):
                catalogue["titles"][title_id]["outfit"] = data["outfit"]
            if data.get("route"):
                catalogue["titles"][title_id]["route"] = data["route"]
            if data.get("age_band"):
                catalogue["titles"][title_id]["age_band"] = data["age_band"]
    CATALOGUE.write_text(json.dumps(catalogue, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("Updated catalogue approved flags.")
    return 0


def find_relative_safe(art: str) -> bool:
    """True if art looks safe to keep (no Same-spot style refs)."""
    from pipeline.catalogue import find_relative_art_refs

    return not find_relative_art_refs(art)


if __name__ == "__main__":
    # Allow importing catalogue relative patterns when run from book_factory
    sys.path.insert(0, str(ROOT))
    raise SystemExit(main())
